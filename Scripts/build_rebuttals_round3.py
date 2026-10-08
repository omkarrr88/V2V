"""
build_rebuttals_round3.py
=========================
Builds paper/rebuttals/Reviewer_Comments_Consolidated.docx for ROUND 3:
the verbatim Round-3 comments from Reviewer 1, each followed by a
point-by-point author response describing the exact manuscript changes.

(The Round-1 and Round-2 consolidated documents are archived at
 paper/rebuttals/round1_archive/Reviewer_Comments_Consolidated_round1.docx
 paper/rebuttals/round2_archive/Reviewer_Comments_Consolidated_round2.docx)
"""
from pathlib import Path

from docx import Document
from docx.shared import RGBColor, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

PROJ = Path(__file__).resolve().parent.parent
OUT  = PROJ / 'paper' / 'rebuttals' / 'Reviewer_Comments_Consolidated.docx'

BLUE = RGBColor(0x1F, 0x3F, 0x7F)
GREY = RGBColor(0x55, 0x55, 0x55)

# ---------------------------------------------------------------------------
# Verbatim Round-3 reviewer comments
# ---------------------------------------------------------------------------
COMMENTS = {
    'Reviewer 1': """Round 3 comments:

1. Equations (4) and (13) are not clearly presented in manuscript.

2. In addition, since the Results section contains multiple subsections, it is recommended to provide an analysis framework or flowchart to clearly illustrate the function of each subsection and the logical relationships among them.""",
}

# ---------------------------------------------------------------------------
# Author responses (Round 3)
# ---------------------------------------------------------------------------
RESPONSES = {
'Reviewer 1': [

("Comment R1.1 - Equations (4) and (13) are not clearly presented",
 "Equations (4) and (13) are not clearly presented in manuscript.",

 "We investigated this carefully and found two distinct causes, both of which "
 "have been fixed.\n\n"

 "(a) Typesetting of the two equations. Equations (4) and (13) were the two "
 "most deeply nested displays in the paper, and both were written as a single "
 "wide line that is cramped in a journal column. Equation (4) placed a fraction "
 "inside a three-argument clamp() call on one line, and the operator itself was "
 "written as \"clamp\" here but as clip with a [0,1] subscript elsewhere in the "
 "paper, so the same operation appeared under two different names. "
 "Equation (13) nested a squared fraction inside a three-case piecewise "
 "definition whose conditions were long enough to crowd the column.\n\n"

 "Both are now presented in two short steps instead of one wide line, using a "
 "named intermediate quantity. Equation (4) first defines a dimensionless speed "
 "fraction s(v_e) as the clipped normalised speed, then writes the zone length "
 "as a simple linear interpolation L_bs = L_base + lambda_scale * s(v_e); this "
 "also now matches the implementation exactly, which computes the clipped "
 "fraction into a temporary variable before scaling it. Equation (13) first "
 "defines a dimensionless urgency ratio kappa = TTC_crit / TTC together with its "
 "cut-off kappa_0 = TTC_crit / TTC_max, then states the three regimes as "
 "kappa >= 1, kappa_0 <= kappa < 1, and kappa < kappa_0, which reduces each case "
 "condition from a long double inequality in TTC to a short one in kappa. Both "
 "restructured forms are mathematically identical to the originals at every "
 "point, including both breakpoints, the TTC = infinity case and the TTC = 0 "
 "case, and both match the reference implementation branch for branch.\n\n"

 "Each equation is now followed by a short paragraph that reads the formula line "
 "by line, states what every symbol means, and gives worked numeric values. For "
 "Eq. (4): v_e = 1.0 m/s gives s = 0 and L_bs = 4.5 m, v_e = 13.0 m/s gives "
 "s = 0.29 and L_bs = 8.0 m, v_e >= 40.0 m/s gives s = 1 and L_bs = 16.5 m. For "
 "Eq. (13), at the published thresholds: TTC = 2 s and 4 s give 1, TTC = 6 s "
 "gives 4/9, TTC = 8 s gives 0.25, and anything beyond 8 s gives 0; the "
 "paragraph also states where the component is continuous and where the drop to "
 "zero at TTC_max is deliberate.\n\n"

 "Reworking the two equations prompted a consistency audit of the notation "
 "around them, which produced four further corrections. The inconsistent "
 "clamp() notation is gone: the unit-interval bound is defined once, at "
 "Eq. (4), and reused by name in Eqs. (11) and (16), which also let us delete "
 "the duplicate definition Eq. (11) previously carried and the undefined "
 "three-argument clip() that the CRI formula (16) still used. Equation (13) now "
 "names its left-hand side R_ttc,lon rather than R_ttc, because Eq. (14) "
 "combines it with a lateral branch; that lateral branch, previously described "
 "only in running prose, is now written out as a displayed piecewise with its "
 "domain and its two degenerate cases stated. The six zone-limit symbols in "
 "Eqs. (5)-(6) - x_min, x_max, x_c, y_min, y_max, y_c - were previously used "
 "without definition and are now defined explicitly, which also connects "
 "Eq. (4) to the rest of the model by showing that L_bs enters the risk score "
 "as the lower longitudinal limit y_min. The notation table (Table 1) gained "
 "rows for the clip operator, s, the ego dimensions and lane width, the GPS "
 "standard deviation, the friction coefficient and air density, kappa and "
 "kappa_0, and the two TTC branches.\n\n"

 "(b) Rendering of multi-line equations in the Word conversion. While checking "
 "the two equations we found a defect in the LaTeX-to-Word conversion itself, "
 "which affected every multi-line display in the manuscript and was almost "
 "certainly the more visible half of the problem. The converter was collapsing "
 "the rows of each multi-line equation onto a single line and emitting the "
 "LaTeX column-alignment markers as literal ampersand characters, so an "
 "equation that is correct in the typeset source appeared in the Word file as, "
 "for example, \"CRI &= clip(...)\" run together on one line. The conversion "
 "pipeline has been fixed: the submitted Word file now contains zero stray "
 "alignment characters, and every multi-line equation - including (2)-(3), (4), "
 "(13) and (16) - is a native Word equation with its rows correctly stacked.\n\n"

 "The restructuring deliberately adds and removes no numbered display, so the "
 "manuscript still has exactly 16 numbered equations and the two the reviewer "
 "identified are still numbered (4) and (13). The one display added in this "
 "round, the lateral TTC branch, is intentionally left unnumbered for that "
 "reason.",

 "Section 4.2 (Eq. (4) restructured, new reading paragraph with worked values); "
 "Section 4.3 (the six zone-limit symbols of Eqs. (5)-(6) defined); "
 "Section 4.5 (Eq. (11) simplified, duplicate clip definition removed); "
 "Section 4.6 (Eq. (13) restructured and renamed R_ttc,lon, new reading "
 "paragraph, lateral branch written out, Eq. (14) tightened); "
 "Section 4.8 (Eq. (16) now uses the defined clip operator); "
 "Table 1 extended with seven new symbol rows and two revised ones; roadmap "
 "paragraph at the head of "
 "Section 4 states the two-step presentation convention; build_paper_docx.py "
 "fixed so multi-line equations export as native Word equations. Equation "
 "numbering is unchanged."),

("Comment R1.2 - Analysis framework / flowchart for the Results section",
 "In addition, since the Results section contains multiple subsections, it is "
 "recommended to provide an analysis framework or flowchart to clearly illustrate "
 "the function of each subsection and the logical relationships among them.",

 "We agree, and have added both a flowchart and a framework table at the head of "
 "Section 7, before the first results subsection.\n\n"

 "The ten results subsections are now grouped into five stages: Stage 1, "
 "validity of the measuring stick (7.1, ground-truth definition); Stage 2, how "
 "well the engine detects (7.2 overall performance, 7.3 calibration discipline, "
 "7.4 ROC and PR analysis); Stage 3, whether a learned model would be a better "
 "choice (7.5 exploratory AI model, 7.6 comparison with other ML methods); "
 "Stage 4, whether the result survives a change of conditions (7.7 per-scenario, "
 "7.8 multi-seed); and Stage 5, why it works and where it breaks (7.9 ablation, "
 "7.10 sensitivity). A short paragraph states what each stage contributes and "
 "how it relates to the others.\n\n"

 "We were careful to draw the actual logical structure rather than the reading "
 "order, because the two are not the same here. Stages 1 and 2 are genuinely "
 "sequential - nothing can be measured before the ground truth is defined - but "
 "Stages 3, 4 and 5 are three independent interrogations of the Stage-2 result "
 "and depend on each other not at all. The new Fig. 8 therefore draws a hub, "
 "not a chain: Stage 1 feeds Stage 2, and the Stage-2 result then fans out to "
 "Stages 3, 4 and 5, which may be read in any order. Each stage is one "
 "colour-coded band carrying a card per subsection, each card giving the "
 "subsection number, its title and a one-line note on what it does. Two dashed "
 "arrows, annotated as "
 "notes 1 and 2 under the diagram, mark the two places where a later subsection "
 "sends a conclusion back to an earlier one: the rare-event ceiling established "
 "in Section 7.6 is what keeps the physics engine the sole decision-maker "
 "behind Section 7.2, and the ablation of Section 7.9 is where the published "
 "weights are chosen, which Section 7.3 then re-checks on a held-out partition. "
 "The figure uses the same high-contrast, greyscale-safe palette adopted for "
 "the other figures in the previous round, and its canvas is sized to the final "
 "printed width so that no lettering is reduced below 6 pt on the page.\n\n"

 "The new Table 3 gives the same framework in tabular form, with one row per "
 "subsection and five columns: the stage it belongs to, the subsection itself, "
 "the question it answers, the evidence it presents (the specific tables and "
 "figures), and the outcome it establishes, with the headline number quoted "
 "wherever there is one. A reader can therefore see "
 "the function of each subsection either graphically or as a one-line summary "
 "before reading the section, and a closing paragraph states the five stages as "
 "a single connected argument.\n\n"

 "Please note that inserting one figure and one table at the head of Section 7 "
 "shifts the subsequent numbering: the figures previously numbered 8-18 are now "
 "9-19, and the tables previously numbered 3-10 are now 4-11. Section and "
 "equation numbering are unchanged, and the framework material is deliberately "
 "placed in an unnumbered preamble so that the ten results subsections keep the "
 "numbers 7.1 to 7.10.",

 "New unnumbered 'Analysis framework' preamble at the head of Section 7, with a "
 "five-stage walkthrough and a closing paragraph on the overall argument; new "
 "Fig. 8 (results analysis framework flowchart, generated by "
 "Scripts/generate_results_framework_figure.py); new Table 3 (function of each "
 "results subsection); six new subsection labels added so the figure, the table "
 "and the text cross-reference the subsections directly."),
],
}

PREAMBLE = (
    "This document contains the Round-3 reviewer comments and the authors' "
    "point-by-point responses. Round 3 raised two points, both from Reviewer 1, "
    "and the revision is correspondingly narrow: the presentation of Eqs. (4) and "
    "(13) together with the notation around them and the Word export of every "
    "multi-line equation, and a new analysis framework for the Results section. "
    "No result, number, dataset or conclusion in the paper has changed in this "
    "round. "
    "Two numbering notes for the editor: equation numbering is unchanged, so "
    "Eqs. (4) and (13) are still numbered (4) and (13); but the new Fig. 8 and "
    "new Table 3 at the head of Section 7 shift the following figures from 8-18 "
    "to 9-19 and the following tables from 3-10 to 4-11. "
    "In the revised manuscript with tracked changes "
    "(V2V_paper 08-10-26_changes_red.docx), red marks every paragraph, table "
    "row and equation touched in Round 3, measured against the Round-2 "
    "manuscript; marking is at paragraph granularity, so a paragraph that "
    "gained a sentence is shown red in full. Equation "
    "displays are shown in red only where the mathematics itself changed - "
    "Eqs. (4), (11), (13), (14), (16) and the unnumbered lateral-TTC branch "
    "beside Eq. (14); the other equations are unchanged in content even "
    "though their Word encoding improved. The Round-1 and Round-2 response "
    "documents are archived under paper/rebuttals/round1_archive/ and "
    "paper/rebuttals/round2_archive/."
)


def add_heading(doc, text, size, colour=None, bold=True, centre=False):
    p = doc.add_paragraph()
    if centre:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if colour is not None:
        r.font.color.rgb = colour
    return p


def main():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)

    add_heading(doc, 'Reviewer Comments and Author Responses - Round 3 (Consolidated)',
                16, centre=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Manuscript: Vehicle-to-Vehicle Communication for Blind Spot '
                  'Detection and Accident Prevention Using Severity-Gated '
                  'Collision Risk Indexing')
    r.italic = True
    r.font.size = Pt(10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Corresponding author: Deepali Vora (deepali.vora11@sitpune.edu.in)'
                  '   |   8 October 2026')
    r.font.size = Pt(10)
    r.font.color.rgb = GREY
    doc.add_paragraph()

    add_heading(doc, 'Scope of this revision', 12)
    doc.add_paragraph(PREAMBLE)
    doc.add_paragraph()

    for reviewer in ['Reviewer 1']:
        add_heading(doc, reviewer, 15, colour=BLUE)

        add_heading(doc, 'Original comments (Round 3, verbatim):', 12)
        for block in COMMENTS[reviewer].strip().split('\n\n'):
            doc.add_paragraph(block.strip())
        doc.add_paragraph()

        add_heading(doc, 'Author responses:', 12)
        for title, quoted, response, change in RESPONSES[reviewer]:
            add_heading(doc, title, 11.5)
            p = doc.add_paragraph()
            r = p.add_run('Reviewer comment: ')
            r.bold = True
            r = p.add_run(quoted)
            r.italic = True
            for i, block in enumerate(response.split('\n\n')):
                p = doc.add_paragraph()
                if i == 0:
                    r = p.add_run('Response: ')
                    r.bold = True
                p.add_run(block.strip())
            p = doc.add_paragraph()
            r = p.add_run('Manuscript change: ')
            r.bold = True
            p.add_run(change)
            doc.add_paragraph()

        sep = doc.add_paragraph('-' * 60)
        sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph()

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f'wrote: {OUT}')


if __name__ == '__main__':
    main()
