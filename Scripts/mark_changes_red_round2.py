"""
mark_changes_red_round2.py
==========================
Round-2 red-marking: colour RED exactly the content changed in the round-2
revision, using the round-1 docx as baseline.

Because the round-1 docx and the round-2 docx encode inline math differently
(unicode text vs OMML), naive fuzzy text comparison produces false positives
near equations. Body paragraphs are therefore marked from an explicit
whitelist of round-2 change fragments (every edit made for Round 2 is known),
and a fuzzy pass merely WARNS about non-whitelisted paragraphs that look
changed so omissions can be reviewed. Table cells are compared per-table
against the best-matching round-1 table; entirely new tables go fully red.

Baseline : paper/rebuttals/round1_archive/V2V_paper 24-06-26_round1.docx
Input    : paper/V2V_paper 14-08-26.docx
Output   : paper/V2V_paper 14-08-26_changes_red.docx
"""
import re
from difflib import SequenceMatcher
from pathlib import Path

from docx import Document
from docx.shared import RGBColor

PROJ     = Path(__file__).resolve().parent.parent
DOCX_OLD = PROJ / 'paper' / 'rebuttals' / 'round1_archive' / 'V2V_paper 24-06-26_round1.docx'
DOCX_IN  = PROJ / 'paper' / 'V2V_paper 14-08-26.docx'
DOCX_OUT = PROJ / 'paper' / 'V2V_paper 14-08-26_changes_red.docx'

RED = RGBColor(0xC0, 0x00, 0x00)

# ---------------------------------------------------------------------------
# Round-2 change whitelist. A body paragraph is coloured red iff its
# normalised text contains one of these fragments. Each fragment corresponds
# to a specific Round-2 edit of main.tex.
# ---------------------------------------------------------------------------
CHANGE_FRAGMENTS = [
    # Abstract (updated Tier-1 numbers + baseline wording)
    'statistically significantly above the ttc kinematic',
    # Intro scope statement (Reviewer 2.2)
    'we state the scope of this study explicitly at the outset',
    # Fig. 1 caption with movement sequence (Reviewer 1.1)
    'driving directions and movement sequence',
    # Related-work + TTC-section Hayward/NHTSA citations (Reviewer 1.3)
    'introduced by hayward',
    'nhtsa forward collision warning band',
    'the thresholds are not free hyperparameters',
    'real-world wireless signals can be interrupted',
    # Mathematical-model roadmap + notation table (Reviewer 1.2)
    'the model is presented bottom-up in seven steps',
    'notation used in the mathematical model',
    'combines the zone-occupancy probability',
    # Tier-1 independence qualifications (Reviewer 2.1)
    'two qualifications keep this claim honest',
    # Dataset integrity audit note (Reviewer 4)
    'automated integrity audit',
    # Bootstrap citation + updated overall-performance prose
    'all auc values carry 95',
    'the physics model is statistically better than every baseline on both',
    # Calibration section (updated numbers)
    'the published weights generalise to the held-out test partition at proxy',
    # ROC / PR purpose statements (Reviewer 1.5) + updated numbers
    'compare threshold-independent detection quality',
    'what fraction of raised alerts would actually correspond',
    # Exploratory AI results (updated numbers + argmax wording)
    'included as an exploratory learning-based baseline',
    'where the critical mass is absorbed',
    # Gain-based terminology (Reviewer 1.4)
    'gain-based split contribution',
    'gain-based split contributions',
    'identify redundant inputs',
    'the gain-based split statistic reported',
    # SHAP purposes + sample + validity (Reviewer 1.5, 2.3)
    'serve distinct functions',
    'a fair question is whether shap output is meaningful',
    'decel first followed by max_gap.',
    # Other-ML section rewrite + operating points (Reviewer 2.3)
    'the three additional classifiers all sit in the same performance band',
    'cost-sensitive operating points',
    'because a 0% recall figure invites the objection',
    'cost-sensitive critical operating points on the held-out test',
    'precision-recall curves for the binary task critical',
    # Speed-CRI + calibration figure purposes (Reviewer 1.5)
    'plots ego speed against the per-step maximum cri',
    # Ablation / multiseed rounding cross-references
    'the same quantity reported rounded as',
    # Discussion updates
    'model-independent ranking physics',
    'even granting each learner a tuned decision threshold',
    'that check itself rests on only 27 collision events',
    # Conclusion updated numbers
    'against a label-independent collision ground truth',
    'hit the same rare-event ceiling as xgboost',
    # Post-audit fix batch (round-2 second pass)
    'organized as a four-layer pipeline',
    'early v2v systems used shared gps coordinates',
    'the speed normalisation runs from',
    'the longitudinal gap evolves as',
    'we deliberately do not reuse the ego-frame coordinate symbols',
    'two training-only augmentations are applied',
    'is label-independent',
    'statistically indistinguishable',
    'the opposite of what an input-overlap artefact would produce',
    'reports two operating points per classifier',
    'f1 at the 0.80 threshold is small for every configuration',
    'mean cri per simulation step by scenario type',
    'hnr sustains the highest mean cri',
    'carries the largest mean',
    'establish the operating envelope',
    'resamples of the full',
    # Verification pass-2 fixes
    'is the signed net acceleration recovered',
    'under the default argmax decision rule',
    'elevated-cri points occur across the whole ego-speed range',
    # Text present in round-2 tex but absent from the round-1 docx the
    # reviewers actually received (late round-1 edits that never reached it)
    'current status all results in this paper are from simulation',
    'the physics model is the sole decision-maker',
    # New bibliography entries (Reviewer 1.3)
    'near-miss determination',
    'extended time-to-collision measures',
    'forward collision warning system confirmation test',
    'estimates of error rates for codes',
    'an introduction to roc analysis',
    'the relationship between precision-recall',
    'an introduction to the bootstrap',
]


def normalise(s: str) -> str:
    s = re.sub(r'[^\w.\[\]%-]', ' ', s.lower())
    s = re.sub(r'\s+', ' ', s)
    return s.strip()


def norm_ascii(s: str) -> str:
    """ASCII-only normalisation for table cells: the round-1 build embeds
    math as unicode text while the round-2 build uses OMML (invisible to
    text extraction), so only the ASCII part of a cell is comparable."""
    return normalise(s.encode('ascii', 'ignore').decode())


from docx.oxml.ns import qn
M_NS = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'


def _colour_omml(paragraph):
    """Colour OMML math runs red (python-docx run API does not reach them)."""
    for mr in paragraph._p.iter(M_NS + 'r'):
        wrPr = mr.find(qn('w:rPr'))
        if wrPr is None:
            wrPr = mr.makeelement(qn('w:rPr'), {})
            mrPr = mr.find(M_NS + 'rPr')
            if mrPr is not None:
                mrPr.addnext(wrPr)
            else:
                mr.insert(0, wrPr)
        color = wrPr.find(qn('w:color'))
        if color is None:
            color = wrPr.makeelement(qn('w:color'), {})
            wrPr.append(color)
        color.set(qn('w:val'), 'C00000')


def has_omml(paragraph) -> bool:
    return next(paragraph._p.iter(M_NS + 'r'), None) is not None


def colour_red(paragraph):
    for run in paragraph.runs:
        if run.text.strip():
            run.font.color.rgb = RED
    _colour_omml(paragraph)


def old_table_rows(doc):
    """Old tables as row lists of ascii-normalised cells + flat cell sets."""
    out = []
    for tbl in doc.tables:
        rows, cells = [], set()
        for row in tbl.rows:
            r = [norm_ascii(cell.text) for cell in row.cells]
            rows.append(r)
            cells.update(t for t in r if t)
        out.append({'rows': rows, 'cells': cells})
    return out


def best_match_ratio(new_norm, old_paras):
    best = 0.0
    for op in old_paras:
        if not op:
            continue
        if max(len(op), len(new_norm)) > 10 * min(len(op), len(new_norm)) + 5:
            continue
        if new_norm in op or op in new_norm:
            return 1.0
        r = SequenceMatcher(None, new_norm, op, autojunk=False).ratio()
        best = max(best, r)
        if best > 0.95:
            return best
    return best


def main():
    print(f'Baseline (round 1): {DOCX_OLD.name}')
    old_doc = Document(str(DOCX_OLD))
    old_paras = [normalise(p.text) for p in old_doc.paragraphs if p.text.strip()]
    old_tables = old_table_rows(old_doc)

    doc = Document(str(DOCX_IN))
    changed = unchanged = 0
    warnings = []

    # ---- Body paragraphs: whitelist ----
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        norm = normalise(txt)
        if any(frag in norm for frag in CHANGE_FRAGMENTS):
            colour_red(p)
            changed += 1
        else:
            unchanged += 1
            # advisory fuzzy check for possible missed changes
            if len(norm.split()) >= 6:
                if best_match_ratio(norm, old_paras) < 0.60:
                    warnings.append(txt[:110])

    # ---- Tables: per-table comparison with explicit policies ----
    for t_i, tbl in enumerate(doc.tables):
        new_rows = [[norm_ascii(c.text) for c in row.cells] for row in tbl.rows]
        new_cells = {t for r in new_rows for t in r if t}
        # explicit policies by table signature
        if 'symbol' in new_cells and 'meaning' in new_cells:
            policy = 'new'          # notation table (added in round 2)
        elif 'screening recall' in new_cells:
            policy = 'new'          # operating-point table (added in round 2)
        elif 'config' in new_cells and 'lat ttc' in new_cells:
            policy = 'ablation'     # F1(0.80) column corrected in round 2
        else:
            policy = 'compare'
        # best-matching old table by cell overlap
        best, best_ov = None, 0.0
        if policy == 'compare':
            for ot in old_tables:
                if not ot['cells']:
                    continue
                ov = len(new_cells & ot['cells']) / max(1, len(new_cells))
                if ov > best_ov:
                    best_ov, best = ov, ot
        # row alignment by first-cell label (positional cell comparison)
        aligned = {}
        if best is not None:
            old_by_label = {r[0]: r for r in best['rows'] if r and r[0]}
            for r_i, nr in enumerate(new_rows):
                if nr and nr[0] and nr[0] in old_by_label \
                        and len(old_by_label[nr[0]]) == len(nr):
                    aligned[r_i] = old_by_label[nr[0]]
        ABLATION_RED = {'0.029', '0.024', '0.001'}
        # Cells verified unchanged since round 1 whose old-encoding text
        # defeats both alignment and set membership:
        NEVER_RED = {'decel. weight'}
        for r_i, row in enumerate(tbl.rows):
            for c_i, cell in enumerate(row.cells):
                t = norm_ascii(cell.text)
                if policy == 'new':
                    red = True
                elif policy == 'ablation':
                    red = t in ABLATION_RED
                elif not t:
                    continue
                elif t in NEVER_RED:
                    red = False
                elif r_i in aligned:
                    red = aligned[r_i][c_i] != t
                else:
                    red = best is None or t not in best['cells']
                for p in cell.paragraphs:
                    if not (p.text.strip() or has_omml(p)):
                        continue
                    if red:
                        colour_red(p)
                        changed += 1
                    else:
                        unchanged += 1
        tag = policy if policy != 'compare' else \
            f'overlap {best_ov:.0%}, {len(aligned)}/{len(new_rows)} rows aligned'
        print(f'  table {t_i}: {tag}')

    print(f'  RED: {changed}   black: {unchanged}')
    if warnings:
        print('  ADVISORY - non-whitelisted paragraphs with low similarity to round 1:')
        for w in warnings:
            print('   ?', w)
    doc.save(str(DOCX_OUT))
    print(f'  wrote: {DOCX_OUT}')

    # ---- verification: text & structure identical ----
    a, b = Document(str(DOCX_IN)), Document(str(DOCX_OUT))
    ok = (len(a.paragraphs) == len(b.paragraphs)
          and len(a.tables) == len(b.tables)
          and all(pa.text == pb.text for pa, pb in zip(a.paragraphs, b.paragraphs)))
    print('  VERIFY OK' if ok else '  VERIFY FAILED')


if __name__ == '__main__':
    main()
